"""Structural checks on the workbook, runnable without Excel or WPS.

The README makes four countable promises about templates/*.xlsx: five sheets,
23 requirement rows, fourteen columns, three dropdown-constrained columns.
All four were checked by hand before; this makes them fail loudly instead.

Usage: python tests/xlsx.py
"""
import glob
import os
import re
import sys

try:
    import openpyxl
except ImportError:
    print('SKIP  openpyxl not installed (pip install openpyxl)')
    sys.exit(0)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHEETS = ['0-说明', '1-GSPR对照表', '2-自查统计', '3-填写示例', '4-签署与版本']
EXPECTED_COLS = 14
EXPECTED_DROPDOWNS = 3

fails = []
checks = 0


def check(label, ok, detail=''):
    global checks
    checks += 1
    if ok:
        print('  ok    ' + label)
    else:
        fails.append(label)
        print('  FAIL  ' + label + (('  —  ' + detail) if detail else ''))


def main():
    files = glob.glob(os.path.join(ROOT, 'templates', '*.xlsx'))
    check('a workbook exists in templates/', len(files) == 1, 'found %d' % len(files))
    if not files:
        return

    p = files[0]
    wb = openpyxl.load_workbook(p)
    print('  file: %s' % os.path.basename(p))

    check('five sheets, in the documented order', wb.sheetnames == SHEETS,
          'got %s' % wb.sheetnames)

    ws = wb['1-GSPR对照表']
    header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    ncols = len([h for h in header if h])
    check('14 columns (README says fourteen)', ncols == EXPECTED_COLS, 'got %d' % ncols)

    # 23 requirement numbers in column A, no gaps, no duplicates
    nums = []
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, (int, float)):
            nums.append(int(v))
    check('23 requirement rows, numbered 1..23', sorted(nums) == list(range(1, 24)),
          'got %s' % nums)

    # Chapter split must be 9 / 13 / 1
    ch = {}
    for r in range(2, ws.max_row + 1):
        c = ws.cell(r, 2).value
        if c:
            ch[str(c).strip()] = ch.get(str(c).strip(), 0) + 1
    check('chapter split is 9 / 13 / 1', list(ch.values()) == [9, 13, 1], 'got %s' % ch)

    # Chapter I carries no official headings — column D must be empty for 1..9
    d_empty = all(not ws.cell(1 + n, 4).value for n in range(1, 10))
    check('points 1-9 leave the official-heading column empty', d_empty,
          'column D is populated for chapter I rows')

    dv = ws.data_validations.dataValidation if ws.data_validations else []
    check('three dropdown-constrained columns', len(dv) == EXPECTED_DROPDOWNS,
          'got %d' % len(dv))
    check('every dropdown rejects out-of-list input (showErrorMessage)',
          all(d.showErrorMessage for d in dv),
          'a validation range exists that silently accepts anything')

    # The point-9 reminder must cover the whole row, not one cell.
    p9 = [(str(list(rng.sqref.ranges)[0]), list(r.formula or []))
          for rng, rules in ws.conditional_formatting._cf_rules.items()
          for r in rules if any('E10' in str(x) for x in (r.formula or []))]
    check('the point-9 reminder covers the whole row, not a single cell',
          bool(p9) and p9[0][0] == 'A10:N10', 'ranges: %s' % p9)

    # Self-check page must never contain a verdict word
    ws2 = wb['2-自查统计']
    txt = ' '.join(str(ws2.cell(r, c).value or '')
                   for r in range(1, ws2.max_row + 1) for c in range(1, 6))
    check('self-check page carries no "ready to submit" verdict',
          not re.search(r'可以(送审|提交)|ready to (submit|file)', txt),
          'found a verdict phrase in the self-check sheet')
    check('self-check page states it does not judge',
          '不做判定' in txt, 'missing the disclaimer')

    # Sub-point counts quoted in the sheets must match the regulation. 23.1 has
    # (a)-(h) in the OJ text; it was written as (a)-(k) in two cells once, in
    # the same file, and nothing caught it because prose is not a test.
    every = ' '.join(str(c.value or '')
                     for sn in wb.sheetnames
                     for row in wb[sn].iter_rows() for c in row)
    check('no sheet claims 23.1 runs to (a)-(k)', '(a)–(k)' not in every,
          'the OJ text of Annex I point 23.1 ends at (h)')
    check('the sub-point ranges quoted for 14 / 18 / 23 are consistent',
          ('14.1–14.7' in every or '14.1-14.7' in every) and
          ('18.1–18.8' in every or '18.1-18.8' in every) and
          ('23.1–23.4' in every or '23.1-23.4' in every),
          'a quoted sub-point range is missing or spelled differently')


main()
print('')
if fails:
    print('FAIL  %d of %d checks' % (len(fails), checks))
    sys.exit(1)
print('PASS  %d checks' % checks)
